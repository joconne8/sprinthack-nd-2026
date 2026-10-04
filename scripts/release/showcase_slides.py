"""Build the Aimsigh HTML deck and editable PowerPoint from verified UI captures."""
import argparse
import base64
import html
import hashlib
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "presentation/aimsigh"
ESC = html.escape

CSS = """
*{box-sizing:border-box}body{margin:0;background:#1c2521;color:#1c2521;font-family:Arial,Helvetica,sans-serif}#frame{position:relative;margin:20px auto}#deck{position:relative;width:1600px;height:900px;transform-origin:top left}.slide{position:absolute;inset:0;display:none;overflow:hidden;background:#f6f5ef;padding:68px 88px}.slide.active{display:block}.eyebrow{font-size:19px;letter-spacing:2px;color:#2e7d5b;font-weight:700;margin-bottom:30px}h1{font-family:Georgia,serif;font-size:76px;font-weight:400;line-height:1.08;letter-spacing:-2px;margin:0 0 28px;max-width:1320px}.lead{font-size:29px;line-height:1.45;color:#5d655f;max-width:1280px;margin:0}.cards{display:grid;grid-template-columns:repeat(3,1fr);gap:24px;margin-top:42px}.card{background:#fffef9;border:1px solid #d4ddd3;border-radius:10px;padding:30px;min-height:250px}.card-number{color:#2e7d5b;font-size:18px}h2{font-family:Georgia,serif;font-weight:400;font-size:35px;line-height:1.1;margin:19px 0 18px}.card p{font-size:24px;line-height:1.45;color:#5d655f;margin:0}footer{position:absolute;bottom:28px;left:88px;right:88px;display:flex;justify-content:space-between;align-items:end;gap:24px;font-size:15px;color:#5d655f;border-top:1px solid #ccd5ca;padding-top:15px}footer a{color:inherit}footer .sources{display:flex;gap:15px;flex-wrap:wrap;max-width:1150px}.title h1{font-size:103px;margin-top:80px;max-width:1370px}.title .lead{max-width:1030px;font-size:32px}.signature{display:flex;gap:22px;align-items:center;margin-top:66px;font-size:22px;letter-spacing:2px;color:#2e7d5b}.signature b{font-size:32px}.closing h1{font-size:99px;margin-top:32px}.closing .lead{max-width:1100px}.closing .signature{margin-top:32px}.flow .card{background:#e5eee2;border-color:#c7d8c4}.image h1{font-size:61px;width:455px;max-width:455px}.image .lead{width:435px;font-size:27px;margin-top:36px}.product-frame{position:absolute;top:113px;right:68px;width:925px;height:646px;border:1px solid #cbd6c8;background:white;border-radius:12px;overflow:hidden;box-shadow:0 16px 46px #1c252119}.window-bar{height:42px;background:#e5ebe2;display:flex;align-items:center;gap:8px;padding:0 18px}.window-bar i{height:9px;width:9px;background:#9daf98;border-radius:50%}.window-bar span{font-size:14px;color:#5d655f;margin-left:10px}.product-frame img{width:100%;height:564px;object-fit:contain;display:block}.image-caption{padding:8px 14px;font-size:15px;line-height:1.3;color:#5d655f}.capture-pending{display:flex;align-items:center;justify-content:center;height:564px;padding:60px;font-size:26px;line-height:1.5;color:#5d655f;background:#fafaf6}.source-grid{display:grid;grid-template-columns:repeat(3,1fr);gap:14px;margin-top:22px}.source-card{padding:14px 22px;border:1px solid #d4ddd3;border-radius:9px;background:#fffef9}.source-card strong{display:block;font-size:24px;margin-bottom:10px}.source-card span{font-size:20px;color:#5d655f;line-height:1.3}.badge{display:inline-block;padding:4px 8px;border-radius:4px;font-size:14px;color:#2e7d5b;background:#e5eee2;margin-bottom:8px}.flow-row{display:flex;gap:13px;align-items:center;margin-top:30px}.flow-box{flex:1;background:#e5eee2;border:1px solid #c7d8c4;padding:23px 14px;border-radius:8px;text-align:center;font-size:24px}.flow-arrow{font-size:27px;color:#2e7d5b}.erd{display:grid;grid-template-columns:repeat(3,1fr);gap:22px;margin-top:25px}.erd-box{background:#fffef9;border:1px solid #d4ddd3;border-radius:8px;padding:23px;font-size:22px;line-height:1.5}.erd-box strong{color:#2e7d5b}.erd-note{font-size:19px;color:#5d655f;margin-top:16px}.controls{display:flex;justify-content:center;align-items:center;gap:13px;color:#e7eee3;font-size:15px;flex-wrap:wrap}button{background:#2e7d5b;color:white;border:1px solid #8baa91;border-radius:6px;padding:9px 15px;cursor:pointer}.controls a{color:#e7eee3}.notes{position:fixed;bottom:0;left:0;right:0;z-index:100;background:#fffef9;color:#1c2521;padding:22px 10%;font-size:19px;line-height:1.5;display:none}.notes.open{display:block}.notes strong{display:block;margin-bottom:8px}.video-modal{position:fixed;inset:0;background:#1c2521ee;z-index:110;display:none;align-items:center;justify-content:center;flex-direction:column;gap:20px}.video-modal.open{display:flex}.video-modal video{width:min(90vw,1280px);max-height:78vh}.video-modal p{color:white;font-size:19px}.draft{position:absolute;right:90px;top:35px;color:#865c18;font-size:16px}.status{font-size:14px;color:#e7eee3;text-align:center;margin-top:12px}
@media print{@page{size:1600px 900px;margin:0}body{background:white}#frame{margin:0;width:1600px!important;height:auto!important}#deck{height:auto;transform:none!important}.slide{position:relative;display:block!important;width:1600px;height:900px;break-after:page;-webkit-print-color-adjust:exact;print-color-adjust:exact}.slide:last-child{break-after:auto}.controls,.notes,.video-modal,.status{display:none!important}}
"""
JS = """
const slides=[...document.querySelectorAll('.slide')];let current=0;
function fit(){const scale=Math.max(.1,Math.min((innerWidth-30)/1600,(innerHeight-105)/900));document.querySelector('#deck').style.transform='scale('+scale+')';document.querySelector('#frame').style.width=1600*scale+'px';document.querySelector('#frame').style.height=900*scale+'px';}
function show(index){current=Math.max(0,Math.min(slides.length-1,index));slides.forEach((s,i)=>s.classList.toggle('active',i===current));document.querySelector('#counter').textContent=(current+1)+' / '+slides.length;document.querySelector('#note-text').textContent=slides[current].dataset.notes;document.querySelector('#note-label').textContent=slides[current].dataset.section+' · '+slides[current].dataset.seconds+' seconds';history.replaceState(null,'','#'+(current+1));}
document.querySelector('#next').onclick=()=>show(current+1);document.querySelector('#prev').onclick=()=>show(current-1);document.querySelector('#notes-toggle').onclick=()=>document.querySelector('.notes').classList.toggle('open');document.querySelector('#fullscreen').onclick=()=>document.fullscreenElement?document.exitFullscreen():document.documentElement.requestFullscreen();
document.querySelector('#video-toggle').onclick=()=>document.querySelector('.video-modal').classList.add('open');document.querySelector('#video-close').onclick=()=>{document.querySelector('.video-modal video').pause();document.querySelector('.video-modal').classList.remove('open')};
document.addEventListener('keydown',e=>{if(e.target.tagName==='VIDEO')return;if(e.key==='Escape'){document.querySelector('#video-close').click();document.querySelector('.notes').classList.remove('open')}if(['ArrowRight','PageDown',' '].includes(e.key)){e.preventDefault();show(current+1)}if(['ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();show(current-1)}if(e.key==='Home')show(0);if(e.key==='End')show(slides.length-1);if(e.key.toLowerCase()==='n')document.querySelector('.notes').classList.toggle('open')});addEventListener('resize',fit);fit();show(Math.max(0,(Number(location.hash.slice(1))||1)-1));
"""

SOURCES = [
    ("Upright", "connected local replica", "Paid-order sales · portal initiation"),
    ("Cash Monkey", "connected local replica", "Unit/order sales · portal CSV export"),
    ("Jewelry", "fixture / access pending", "Provided report · mappings unconfirmed"),
    ("OSM / PB / EasyPost", "fixture / access pending", "Shipping expenses · export or lookup"),
    ("FedEx", "fixture / access pending", "Charges and refunds · expense control"),
    ("ShopGoodwill", "fixture / access pending", "Marketplace report · overlap review"),
    ("Goodwill Books", "fixture / access pending", "Monthly statement · settlement control"),
    ("eBay", "fixture / access pending", "Marketplace report · overlap review"),
    ("Amazon", "fixture / access pending", "Payment summary · settlement control"),
]


def build_html(slides, sources, draft):
    if not draft:
        verification_path = OUT / 'capture-verification.json'
        if not verification_path.is_file():
            raise SystemExit('Run the actual showcase capture/rehearsal before building a final deck')
        verification = json.loads(verification_path.read_text())
        if verification.get('passed') is not True:
            raise SystemExit('The actual application capture has not passed')
        for name, digest in verification.get('capture_sha256', {}).items():
            if hashlib.sha256((OUT / name).read_bytes()).hexdigest() != digest:
                raise SystemExit('Capture changed since verification: ' + name)
    sections = []
    missing = []
    for index, slide in enumerate(slides, 1):
        kind = slide.get("kind", "card-slide")
        parts = [f'<div class="eyebrow">{ESC(slide["eyebrow"])}</div>',
                 f'<h1>{ESC(slide["title"]).replace(chr(10), "<br>")}</h1>',
                 f'<p class="lead">{ESC(slide["body"])}</p>']
        if "cards" in slide:
            cards = [f'<div class="card"><span class="card-number">0{i+1}</span><h2>{ESC(c["title"])}</h2><p>{ESC(c["text"])}</p></div>' for i,c in enumerate(slide["cards"])]
            parts.append('<div class="cards">'+''.join(cards)+'</div>')
        if "image" in slide:
            capture = OUT / slide["image"]
            if capture.is_file():
                encoded = base64.b64encode(capture.read_bytes()).decode()
                content = f'<img alt="{ESC(slide["image_caption"])}" src="data:image/png;base64,{encoded}">'
            else:
                missing.append(slide["image"])
                content = '<div class="capture-pending">Draft: the real application capture is pending. No invented metric screenshot is substituted.</div>'
            parts.append('<div class="product-frame"><div class="window-bar"><i></i><i></i><i></i><span>ACTUAL SYNTHETIC APPLICATION</span></div>'+content+f'<div class="image-caption">{ESC(slide["image_caption"])}</div></div>')
        if kind == "sources":
            parts.append('<div class="source-grid">'+''.join(f'<div class="source-card"><span class="badge">{ESC(status)}</span><strong>{ESC(name)}</strong><span>{ESC(description)}</span></div>' for name,status,description in SOURCES)+'</div>')
        if kind == "dataflow":
            parts.append('<div class="flow-row">'+('<span class="flow-arrow">→</span>'.join(f'<div class="flow-box">{label}</div>' for label in ["Raw CSV", "Typed staging", "Versioned sales", "Published snapshot", "Excel + answers"]))+'</div>')
            parts.append('<div class="erd"><div class="erd-box"><strong>Original report</strong><br>source_files.id / checksum<br>import_batches.file_id<br>staging_rows.batch_id + row_number</div><div class="erd-box"><strong>Published records</strong><br>sale_versions.id / record_key<br>metric_runs.id<br>metric_run_rows.run_id + version_id</div><div class="erd-box"><strong>Showcase snapshot</strong><br>Source publication identifiers<br>Versioned labor / shipping inputs<br>Scope, coverage and checksums</div></div><p class="erd-note">Actual local Python + SQLite implementation. Money is calculated with integer cents and Decimal; inputs are combined after aggregation at their own grains.</p>')
        if kind in ("title", "closing"):
            parts.append('<div class="signature"><span>COLLECT</span><b>→</b><span>VERIFY & DELIVER</span><b>→</b><span>UNDERSTAND</span></div>')
        source_links = ''.join(f'<a href="{ESC(sources[s]["url"], quote=True)}" title="{ESC(sources[s]["title"],quote=True)}">{ESC(s)}</a>' for s in slide["source_ids"])
        parts.append(f'<footer><span class="sources">Sources: {source_links}</span><span>AIMSIGH / {index:02d}</span></footer>')
        if draft:
            parts.append('<div class="draft">DRAFT · capture verification pending</div>')
        sections.append(f'<article class="slide {kind}" id="slide-{index}" aria-label="Slide {index}" data-section="{ESC(slide["section"],quote=True)}" data-seconds="{slide["seconds"]}" data-notes="{ESC(slide["notes"],quote=True)}">'+''.join(parts)+'</article>')
    if missing and not draft:
        raise SystemExit("Required actual application captures are missing: "+", ".join(missing))
    video_name = "walkthrough.mp4" if (OUT / "walkthrough.mp4").is_file() else "../demo.mp4"
    video_label = "Recorded Aimsigh showcase · synthetic records · local model-free channel" if video_name == "walkthrough.mp4" else "Earlier verified P0 recording · Aimsigh recording not yet generated"
    document = '<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Aimsigh · Goodwill reporting showcase</title><style>'+CSS+'</style></head><body><main id="frame"><div id="deck">'+''.join(sections)+'</div></main><nav class="controls" aria-label="Presentation"><button id="prev">← Previous</button><span id="counter"></span><button id="next">Next →</button><button id="notes-toggle">Speaker notes (N)</button><button id="fullscreen">Fullscreen</button><button id="video-toggle">Recorded backup</button><a href="aimsigh-showcase.pdf">PDF</a><a href="aimsigh-showcase.pptx">PowerPoint</a><a href="'+video_name+'">MP4</a><a href="https://github.com/joconne8/sprinthack-nd-2026">Repository</a></nav><div class="status">15-minute presentation · two synthetic source replicas · invented labor/cost model · local Teams-style prototype</div><aside class="notes"><strong id="note-label"></strong><span id="note-text"></span></aside><aside class="video-modal"><p>'+video_label+'</p><video controls preload="metadata" src="'+video_name+'"></video><button id="video-close">Close recording</button></aside><script>'+JS+'</script></body></html>'
    (OUT / "index.html").write_text(document)
    (OUT / "build.json").write_text(json.dumps({"draft":draft,"slide_count":len(slides),"duration_seconds":sum(s["seconds"] for s in slides),"missing_captures":missing,"video":video_name,"synthetic":True,"labor_costs_invented":True,"native_integrations":False},indent=2)+'\n')


def build_powerpoint(slides, sources):
    if not (OUT / 'capture-verification.json').is_file() or json.loads((OUT / 'capture-verification.json').read_text()).get('passed') is not True:
        raise SystemExit('Final PowerPoint requires a passed actual application capture')
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from PIL import Image
    prs = Presentation(); prs.slide_width=Inches(16); prs.slide_height=Inches(9)
    def box(slide, x,y,w,h,text,size=25,color="1c2521",font="Arial",bold=False):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=shape.text_frame; tf.word_wrap=True
        for i,line in enumerate(text.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        return shape
    def panel(slide,x,y,w,h):
        shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string("e5eee2");shape.line.color.rgb=RGBColor.from_string("c7d8c4")
    for i,s in enumerate(slides,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string("f6f5ef")
        box(slide,.88,.65,14,.4,s['eyebrow'],16,"2e7d5b",bold=True)
        is_image='image' in s; is_title=s.get('kind') in ('title','closing')
        title_y=1.5 if is_title else 1.25
        box(slide,.88,title_y,4.5 if is_image else 14,3.6 if is_title else 2.4,s['title'],52 if is_image else (72 if is_title else 62),font="Georgia")
        box(slide,.88,4.15 if is_image else (5.2 if is_title else 3.35),4.3 if is_image else 13.8,2.1,s['body'],23 if is_image else 25,"5d655f")
        if is_image:
            capture=OUT/s['image']
            if not capture.is_file():raise SystemExit("Missing actual image: "+str(capture))
            image=Image.open(capture);iw,ih=image.size;scale=min(9.2/iw,5.75/ih);width=iw*scale;height=ih*scale
            slide.shapes.add_picture(str(capture),Inches(6.12+(9.2-width)/2),Inches(1.35+(5.75-height)/2),width=Inches(width),height=Inches(height))
            box(slide,6.12,7.15,9.2,.5,s['image_caption'],13,"5d655f")
        if 'cards' in s:
            for n,c in enumerate(s['cards']):
                x=.88+n*4.82;panel(slide,x,4.72,4.55,2.65);box(slide,x+.25,4.95,4.02,.6,c['title'],27,font="Georgia");box(slide,x+.25,5.7,4.02,1.43,c['text'],20,"5d655f")
        if s.get('kind')=='sources':
            for n,(name,status,description) in enumerate(SOURCES):
                x=.88+(n%3)*4.82;y=4.13+(n//3)*1.12;panel(slide,x,y,4.55,1.02);box(slide,x+.15,y+.04,4.2,.3,name+' · '+status,13,"2e7d5b",bold=True);box(slide,x+.15,y+.4,4.2,.6,description,17,"5d655f")
        if s.get('kind')=='dataflow':
            box(slide,.88,4.25,14,.6,'Raw CSV → Typed staging → Versioned sales → Published snapshot → Excel + answers',22,"2e7d5b")
            groups=[('Original report','source_files.id / checksum\nimport_batches.file_id\nstaging_rows.batch_id + row_number'),('Published records','sale_versions.id / record_key\nmetric_runs.id\nmetric_run_rows.run_id + version_id'),('Showcase snapshot','Source publication identifiers\nVersioned labor / shipping inputs\nScope, coverage and checksums')]
            for n,(title,body) in enumerate(groups):
                x=.88+n*4.82;panel(slide,x,5.15,4.55,2.22);box(slide,x+.2,5.35,4.12,.5,title,24,font='Georgia');box(slide,x+.2,5.95,4.12,1.2,body,17,'5d655f')
        if is_title:box(slide,.88,7.05,14,.8,'COLLECT  →  VERIFY & DELIVER  →  UNDERSTAND',23,'2e7d5b')
        box(slide,.88,8.42,13,.3,'Sources: '+' · '.join(s['source_ids'])+'   |   AIMSIGH / '+str(i).zfill(2),11,'5d655f')
        slide.notes_slide.notes_text_frame.text=s['notes']+'\n\nSegment: '+s['section']+' · '+str(s['seconds'])+' seconds\n\n'+'\n'.join(sources[key]['title']+'\n'+sources[key]['url'] for key in s['source_ids'])
    prs.core_properties.title='Aimsigh | Goodwill Michiana';prs.core_properties.subject='15-minute local synthetic reporting showcase';prs.save(OUT/'aimsigh-showcase.pptx')


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--allow-draft',action='store_true');parser.add_argument('--pptx',action='store_true');args=parser.parse_args()
    slides=json.loads((ROOT/'planning/pitch/aimsigh-slides.json').read_text());sources=json.loads((ROOT/'planning/pitch/aimsigh-sources.json').read_text());OUT.mkdir(parents=True,exist_ok=True)
    assert len(slides)==15 and sum(s['seconds'] for s in slides)==900
    if args.pptx:build_powerpoint(slides,sources);print('Built editable 15-slide PowerPoint with full speaker notes and sources.')
    else:build_html(slides,sources,args.allow_draft);print('Built Aimsigh HTML deck'+(' (explicit draft; captures pending).' if args.allow_draft else ' with actual application captures.'))


if __name__=='__main__':main()
